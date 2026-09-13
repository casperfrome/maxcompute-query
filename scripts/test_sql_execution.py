"""Offline execution contracts; fake service objects never contact MaxCompute."""
import importlib
import io
import json
import logging
from pathlib import Path
import subprocess
import signal
import sys
import tempfile
import unittest
from unittest.mock import patch

import pandas as pd
from odps.errors import NoPermission, ODPSError, WaitTimeoutError


class Reader:
    def __init__(self, count=3, columns=('value',), batches=None):
        self.count = count
        self.schema = type('Schema', (), {'names': list(columns)})()
        self.columns = list(columns)
        self.preview_calls = []
        self.batch_calls = []
        self.batches = batches
        self.closed = False

    def __enter__(self):
        return self

    def __exit__(self, *args):
        self.closed = True

    def to_pandas(self, **kwargs):
        self.preview_calls.append(kwargs)
        return pd.DataFrame({c: list(range(kwargs['count'])) for c in self.columns})

    def iter_pandas(self, **kwargs):
        self.batch_calls.append(kwargs)
        if self.batches is not None:
            for batch in self.batches:
                if isinstance(batch, BaseException):
                    raise batch
                yield batch
        else:
            for start in range(0, self.count, kwargs['batch_size']):
                yield pd.DataFrame({c: list(range(start, min(start + kwargs['batch_size'], self.count)))
                                    for c in self.columns})


class Instance:
    def __init__(self, reader=None, query='select 1'):
        self.id = 'instance-123'
        self.project = type('Project', (), {'name': 'project_a'})()
        self.reader = reader or Reader()
        self.query = query
        self.events = []
        self.wait_error = None
        self.success = True
        self.open_errors = []
        self.inspect_durable = None

    def get_logview_address(self):
        self.events.append('logview')
        if self.inspect_durable:
            self.inspect_durable()
        return 'https://example.test/log?Signature=private-signature'

    def wait_for_success(self, **kwargs):
        self.events.append(('wait', kwargs))
        if self.wait_error:
            raise self.wait_error

    def is_successful(self, **kwargs):
        self.events.append(('success', kwargs))
        return self.success

    def is_terminated(self, **kwargs):
        return True

    def open_reader(self, **kwargs):
        self.events.append(('open', kwargs))
        if self.open_errors:
            raise self.open_errors.pop(0)
        return self.reader

    def get_sql_query(self):
        return self.query

    def get_task_results(self):
        self.events.append('raw')
        return {'SQLTask': 'raw description'}


class Client:
    project = 'project_a'
    endpoint = 'https://service.test'

    def __init__(self, instance=None):
        self.instance = instance or Instance()
        self.submissions = []
        self.recoveries = []
        self.submit_error = None

    def run_sql(self, sql):
        self.submissions.append(sql)
        if self.submit_error:
            raise self.submit_error
        return self.instance

    def get_instance(self, instance_id, **kwargs):
        self.recoveries.append((instance_id, kwargs))
        return self.instance


class ExecutionTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.root = Path(self.tmp.name)
        self.stderr = io.StringIO()
        try:
            self.backend = importlib.import_module('sql_execution')
        except ModuleNotFoundError:
            self.fail('Required sql_execution backend API has not been implemented')

    def execute(self, client=None, **kwargs):
        kwargs.setdefault('sql', 'select 1')
        return self.backend.execute_query(client or Client(), run_root=self.root / 'runs',
                                          stderr=self.stderr, **kwargs)

    def record(self):
        return json.loads(next((self.root / 'runs').glob('*/state.json')).read_text(encoding='utf-8'))

    def test_new_backend_api_exists(self):
        self.assertTrue(callable(self.backend.execute_query))

    def test_submission_is_once_and_id_durable_before_logview_wait(self):
        client = Client()
        def inspect():
            self.assertEqual(self.record()['instance_id'], client.instance.id)
            self.assertIn(client.instance.id, self.stderr.getvalue())
            self.assertIn('project_a', self.stderr.getvalue())
            self.assertIn('state.json', self.stderr.getvalue())
        client.instance.inspect_durable = inspect
        data, meta = self.execute(client)
        self.assertEqual(client.submissions, ['select 1'])
        self.assertEqual(meta['state'], 'succeeded')
        self.assertEqual(meta['project'], 'project_a')
        self.assertEqual(meta['endpoint'], client.endpoint)
        self.assertEqual(meta['downloaded_rows'], 3)
        self.assertEqual(len(data), 3)
        self.assertEqual(self.record()['sql_sha256'], meta['sql_sha256'])
        self.assertEqual((Path(meta['run_dir']) / 'query.sql').read_text(encoding='utf-8'), 'select 1')
        self.assertIn('private-signature', self.record()['logview'])
        self.assertNotIn('private-signature', self.stderr.getvalue())

    def test_timeout_keeps_id_and_recovery_does_not_submit(self):
        client = Client()
        client.instance.wait_error = WaitTimeoutError('timeout')
        with self.assertRaises(self.backend.ExecutionError) as ctx:
            self.execute(client, wait_timeout=7)
        self.assertEqual(ctx.exception.exit_code, 5)
        self.assertEqual(ctx.exception.meta['state'], 'wait_timeout')
        self.assertIsNone(ctx.exception.meta['restricted'])
        self.assertIn("--instance-id 'instance-123'", str(ctx.exception))
        self.assertIn("--project 'project_a'", str(ctx.exception))
        self.assertIn("& '", str(ctx.exception))
        self.assertEqual(len(client.submissions), 1)
        self.assertFalse(any(isinstance(e, tuple) and e[0] == 'open' for e in client.instance.events))
        client.instance.wait_error = None
        data, meta = self.execute(client, sql=None, instance_id='instance-123')
        self.assertEqual(len(client.submissions), 1)
        self.assertEqual(client.recoveries, [('instance-123', {'project': 'project_a'})])
        self.assertEqual(len(data), 3)

    def test_recovery_wrong_project_or_id_never_opens_reader(self):
        for mismatch in ('project', 'id'):
            with self.subTest(mismatch=mismatch):
                client = Client()
                if mismatch == 'project':
                    client.instance.project = 'other'
                else:
                    client.instance.id = 'other'
                with self.assertRaises(self.backend.ExecutionError) as ctx:
                    self.execute(client, sql=None, instance_id='instance-123')
                self.assertEqual(ctx.exception.exit_code, 4)
                self.assertEqual(client.submissions, [])
                self.assertFalse(any(isinstance(e, tuple) and e[0] == 'open' for e in client.instance.events))

    def test_uncertain_submission_is_not_retried_or_called_failed(self):
        client = Client()
        client.submit_error = ConnectionError('socket lost token=do-not-show')
        with self.assertRaises(self.backend.ExecutionError) as ctx:
            self.execute(client)
        self.assertEqual(ctx.exception.exit_code, 4)
        self.assertEqual(ctx.exception.meta['state'], 'submission_unknown')
        self.assertIsNone(ctx.exception.meta['instance_id'])
        self.assertNotIn('do-not-show', str(ctx.exception))
        self.assertEqual(len(client.submissions), 1)

    def test_wait_network_uncertainty_and_remote_failure(self):
        for error, state in [(ConnectionError('lost'), 'wait_unknown'), (ODPSError('task failed'), 'failed')]:
            client = Client()
            client.instance.wait_error = error
            client.instance.success = False
            with self.subTest(state=state), self.assertRaises(self.backend.ExecutionError) as ctx:
                self.execute(client)
            self.assertEqual(ctx.exception.exit_code, 4)
            self.assertEqual(ctx.exception.meta['state'], state)

    def test_success_is_explicitly_verified_before_reading(self):
        client = Client()
        client.instance.success = False
        with self.assertRaises(self.backend.ExecutionError):
            self.execute(client)
        self.assertTrue(any(isinstance(e, tuple) and e[0] == 'success' for e in client.instance.events))
        self.assertFalse(any(isinstance(e, tuple) and e[0] == 'open' for e in client.instance.events))

    def test_interruption_is_recorded_and_reraised(self):
        client = Client()
        client.instance.wait_error = KeyboardInterrupt()
        with self.assertRaises(KeyboardInterrupt):
            self.execute(client)
        self.assertEqual(self.record()['state'], 'interrupted')
        self.assertEqual(len(client.submissions), 1)

    def test_sdk_swallowed_sigint_is_still_recorded_and_never_reads(self):
        client = Client()
        original = signal.getsignal(signal.SIGINT)
        def sdk_wait(**kwargs):
            try:
                signal.getsignal(signal.SIGINT)(signal.SIGINT, None)
            except KeyboardInterrupt:
                pass
        client.instance.wait_for_success = sdk_wait
        with self.assertRaises(KeyboardInterrupt):
            self.execute(client)
        self.assertEqual(self.record()['state'], 'interrupted')
        self.assertIs(signal.getsignal(signal.SIGINT), original)
        self.assertFalse(any(isinstance(e, tuple) and e[0] == 'open' for e in client.instance.events))

    def test_sdk_sigint_followed_by_running_error_is_still_interrupted(self):
        client = Client()
        def sdk_wait(**kwargs):
            try:
                signal.getsignal(signal.SIGINT)(signal.SIGINT, None)
            except KeyboardInterrupt:
                pass
            raise ODPSError('Running')
        client.instance.wait_for_success = sdk_wait
        with self.assertRaises(KeyboardInterrupt):
            self.execute(client)
        self.assertEqual(self.record()['state'], 'interrupted')

    def test_submission_interrupt_without_id_remains_unknown_and_prints_record(self):
        client = Client()
        client.submit_error = KeyboardInterrupt()
        with self.assertRaises(KeyboardInterrupt):
            self.execute(client)
        self.assertEqual(self.record()['state'], 'submission_unknown')
        self.assertIn('state.json', self.stderr.getvalue())
        self.assertIn('不能自动重新提交', self.stderr.getvalue())
        self.assertEqual(len(client.submissions), 1)

    def test_interrupt_with_id_prints_recovery_command(self):
        client = Client()
        client.instance.wait_error = KeyboardInterrupt()
        with self.assertRaises(KeyboardInterrupt):
            self.execute(client)
        self.assertIn("--instance-id 'instance-123'", self.stderr.getvalue())

    def test_sdk_real_progress_does_not_leak_logview_and_filter_is_restored(self):
        from odps.models import Instance as SDKInstance
        client = Client()
        inst = client.instance
        inst.get_all_task_progresses = lambda: {}
        inst._logview_logged = False
        inst._last_progress_value = 0
        inst._last_progress_time = 0
        inst.wait_for_success = lambda **kwargs: SDKInstance._dump_instance_progress(inst, 0, 100000)
        logger = logging.getLogger('odps.models.instance')
        capture = io.StringIO()
        handler = logging.StreamHandler(capture)
        previous_level, previous_filters = logger.level, list(logger.filters)
        logger.addHandler(handler)
        logger.setLevel(logging.INFO)
        try:
            self.execute(client)
            self.assertNotIn('private-signature', capture.getvalue())
            self.assertEqual(logger.filters, previous_filters)
            logger.info('normal progress restored')
            self.assertIn('normal progress restored', capture.getvalue())
        finally:
            logger.removeHandler(handler)
            logger.setLevel(previous_level)

    def test_scanned_summary_is_optional_and_collected_after_reading(self):
        client = Client()
        def summary(name):
            self.assertTrue(client.instance.reader.closed)
            return {'Inputs': {'p.t': [2, 2048]}}
        client.instance.get_task_names = lambda: ['SQLTask']
        client.instance.get_task_summary = summary
        _, meta = self.execute(client)
        self.assertEqual(meta['scanned'], '2.0KB')
        client.instance.get_task_summary = lambda name: (_ for _ in ()).throw(ConnectionError('optional'))
        _, meta = self.execute(client)
        self.assertIsNone(meta['scanned'])

    def test_submission_error_with_known_id_preserves_recovery(self):
        client = Client()
        client.submit_error = ODPSError('lost response', instance_id='instance-123')
        with self.assertRaises(self.backend.ExecutionError) as ctx:
            self.execute(client)
        self.assertEqual(ctx.exception.meta['instance_id'], 'instance-123')
        self.assertIn("--instance-id 'instance-123'", str(ctx.exception))

    def test_invalid_inputs_rejected_before_submission(self):
        invalid = [{'max_rows': 0}, {'wait_timeout': 0}, {'batch_size': 0}, {'save': 'bad.txt'},
                   {'max_rows': True}, {'sql': 'delete from t'}, {'sql': 'show tables', 'save': 'data.csv'},
                   {'instance_id': 'i'}, {'sql': None}]
        for kwargs in invalid:
            client = Client()
            with self.subTest(kwargs=kwargs), self.assertRaises(ValueError):
                self.execute(client, **kwargs)
            self.assertEqual(client.submissions, [])

    def test_module_import_is_offline_without_sdk_pandas_or_config(self):
        code = """import sys, importlib.abc
class Guard(importlib.abc.MetaPathFinder):
 def find_spec(self, fullname, *args):
  if fullname.split('.')[0] in {'odps', 'pandas', 'openpyxl', 'config'}: raise AssertionError(fullname)
sys.meta_path.insert(0, Guard())
import sql_execution
"""
        result = subprocess.run([sys.executable, '-X', 'utf8', '-B', '-c', code],
                                cwd=Path(__file__).parent, capture_output=True, text=True, encoding='utf-8')
        self.assertEqual(result.returncode, 0, result.stderr)

    def test_million_rows_preview_reads_only_requested_rows_once(self):
        reader = Reader(count=1000000)
        client = Client(Instance(reader))
        frame, meta = self.execute(client, max_rows=17)
        self.assertEqual(reader.preview_calls, [{'count': 17, 'n_process': 1}])
        self.assertEqual(reader.batch_calls, [])
        self.assertEqual(len(frame), 17)
        self.assertEqual((meta['total_rows'], meta['downloaded_rows'], meta['displayed_rows']), (1000000, 17, 17))
        self.assertTrue(meta['truncated'])
        self.assertEqual([e[1] for e in client.instance.events if isinstance(e, tuple) and e[0] == 'open'],
                         [dict(tunnel=True, limit=False, arrow=False, reopen=True)])

    def test_zero_preview_keeps_schema_and_never_calls_sdk_count_zero(self):
        reader = Reader(count=0, columns=('a', 'b'))
        frame, meta = self.execute(Client(Instance(reader)))
        self.assertEqual(list(frame.columns), ['a', 'b'])
        self.assertEqual(len(frame), 0)
        self.assertEqual(reader.preview_calls, [])
        self.assertEqual((meta['total_rows'], meta['downloaded_rows'], meta['truncated']), (0, 0, False))

    def test_unknown_count_preview_is_bounded_and_not_falsely_complete(self):
        reader = Reader(count=None)
        frame, meta = self.execute(Client(Instance(reader)), max_rows=2)
        self.assertEqual(reader.preview_calls, [{'count': 2, 'n_process': 1}])
        self.assertEqual(len(frame), 2)
        self.assertIsNone(meta['total_rows'])
        self.assertIsNone(meta['truncated'])

    def test_permission_fallback_only_preview_and_count_is_not_full_count(self):
        client = Client()
        client.instance.open_errors = [NoPermission('protected')]
        frame, meta = self.execute(client, max_rows=2)
        self.assertEqual(len(frame), 2)
        self.assertTrue(meta['restricted'])
        self.assertIsNone(meta['total_rows'])
        self.assertIn(meta['truncated'], (True, None))
        opens = [e[1] for e in client.instance.events if isinstance(e, tuple) and e[0] == 'open']
        self.assertEqual(opens, [dict(tunnel=True, limit=False, arrow=False, reopen=True),
                                dict(tunnel=True, limit=True, arrow=False, reopen=True)])

    def test_permission_export_and_other_preview_errors_never_fallback(self):
        for error, save in [(NoPermission('protected'), self.root / 'data.csv'), (ConnectionError('lost'), None),
                            (ODPSError('unsupported'), None)]:
            client = Client()
            client.instance.open_errors = [error]
            with self.subTest(error=error, save=save), self.assertRaises(self.backend.ExecutionError) as ctx:
                self.execute(client, save=save)
            self.assertEqual(ctx.exception.exit_code, 6)
            self.assertEqual(len([e for e in client.instance.events if isinstance(e, tuple) and e[0] == 'open']), 1)
            if save:
                self.assertFalse(save.exists())

    def test_csv_multibatch_streams_exactly_once_with_one_header_and_bom(self):
        reader = Reader(count=7, columns=('value', '中文'))
        output = self.root / '结果.csv'
        result, meta = self.execute(Client(Instance(reader)), save=output, batch_size=2)
        self.assertIsNone(result)
        self.assertEqual(reader.preview_calls, [])
        self.assertEqual(reader.batch_calls, [{'batch_size': 2, 'n_process': 1}])
        self.assertEqual(output.read_bytes().count(b'\xef\xbb\xbf'), 1)
        lines = output.read_text(encoding='utf-8-sig').splitlines()
        self.assertEqual(lines, ['value,中文'] + [f'{i},{i}' for i in range(7)])
        self.assertEqual((meta['total_rows'], meta['downloaded_rows'], meta['displayed_rows'], meta['truncated']), (7, 7, 0, False))
        self.assertEqual(meta['output_path'], str(output.resolve()))

    def test_zero_csv_and_xlsx_publish_schema_header(self):
        from openpyxl import load_workbook
        for extension in ('csv', 'xlsx'):
            output = self.root / ('empty.' + extension)
            reader = Reader(count=0, columns=('a', 'b'))
            result, meta = self.execute(Client(Instance(reader)), save=output)
            self.assertIsNone(result)
            self.assertEqual(meta['downloaded_rows'], 0)
            self.assertEqual(reader.preview_calls, [])
            if extension == 'csv':
                self.assertEqual(output.read_text(encoding='utf-8-sig').splitlines(), ['a,b'])
            else:
                book = load_workbook(output, read_only=True)
                self.addCleanup(book.close)
                self.assertEqual(list(book.active.values), [('a', 'b')])

    def test_xlsx_multibatch_preserves_nulls_large_integers_and_literal_formula(self):
        from openpyxl import load_workbook
        reader = Reader(count=4, columns=('value',), batches=[
            pd.DataFrame({'value': [pd.NA, '=not-a-formula']}, dtype=object),
            pd.DataFrame({'value': [1234567890123456789, pd.NaT]}, dtype=object)])
        output = self.root / 'data.xlsx'
        _, meta = self.execute(Client(Instance(reader)), save=output, batch_size=2)
        book = load_workbook(output, read_only=True)
        self.addCleanup(book.close)
        self.assertEqual(list(book.active.iter_rows(max_col=1, values_only=True)), [('value',), (None,), ('=not-a-formula',),
                                                   ('1234567890123456789',), (None,)])
        self.assertEqual(meta['downloaded_rows'], 4)
        self.assertEqual(reader.batch_calls, [{'batch_size': 2, 'n_process': 1}])
        self.assertEqual(reader.preview_calls, [])

    def test_export_rejects_unknown_negative_noninteger_counts_before_writing(self):
        for count in (None, -1, 3.5, True):
            output = self.root / 'existing.csv'
            output.write_text('original', encoding='utf-8')
            with self.subTest(count=count), self.assertRaises(self.backend.ExecutionError) as ctx:
                self.execute(Client(Instance(Reader(count=count))), save=output)
            self.assertEqual(ctx.exception.exit_code, 6)
            self.assertEqual(output.read_text(encoding='utf-8'), 'original')

    def test_export_count_mismatch_and_schema_mismatch_preserve_output(self):
        cases = [(3, [pd.DataFrame({'value': [1, 2]})], 2),
                 (1, [pd.DataFrame({'value': [1, 2]})], 2),
                 (2, [pd.DataFrame({'wrong': [1, 2]})], 2),
                 (2, [pd.DataFrame({'value': [1]}), pd.DataFrame({'value': [2], 'extra': [3]})], 2)]
        for total, batches, downloaded in cases:
            output = self.root / 'existing.csv'
            output.write_text('original', encoding='utf-8')
            with self.subTest(total=total, batches=batches), self.assertRaises(self.backend.ExecutionError) as ctx:
                self.execute(Client(Instance(Reader(count=total, batches=batches))), save=output)
            self.assertEqual(ctx.exception.exit_code, 6)
            self.assertEqual(ctx.exception.meta['downloaded_rows'], downloaded)
            self.assertEqual(output.read_text(encoding='utf-8'), 'original')
            self.assertEqual(sorted(p.name for p in self.root.iterdir()), ['existing.csv', 'runs'])

    def test_reader_error_after_first_batch_and_cancel_preserve_output(self):
        for error in (ConnectionError('lost'), KeyboardInterrupt()):
            output = self.root / 'existing.csv'
            output.write_text('original', encoding='utf-8')
            reader = Reader(count=3, batches=[pd.DataFrame({'value': [1]}), error])
            expected = KeyboardInterrupt if isinstance(error, KeyboardInterrupt) else self.backend.ExecutionError
            with self.subTest(error=error), self.assertRaises(expected):
                self.execute(Client(Instance(reader)), save=output)
            self.assertEqual(output.read_text(encoding='utf-8'), 'original')
            self.assertEqual(sorted(p.name for p in self.root.iterdir()), ['existing.csv', 'runs'])

    def test_atomic_replace_failure_preserves_existing_output(self):
        output = self.root / 'existing.csv'
        output.write_text('original', encoding='utf-8')
        real_replace = self.backend.os.replace
        def replace(source, destination):
            if Path(destination) == output:
                raise PermissionError('output locked')
            return real_replace(source, destination)
        with patch.object(self.backend.os, 'replace', side_effect=replace), self.assertRaises(self.backend.ExecutionError) as ctx:
            self.execute(save=output)
        self.assertEqual(ctx.exception.exit_code, 6)
        self.assertEqual(output.read_text(encoding='utf-8'), 'original')
        self.assertEqual(sorted(p.name for p in self.root.iterdir()), ['existing.csv', 'runs'])

    def test_reader_close_failure_happens_before_output_publication(self):
        output = self.root / 'existing.csv'
        output.write_text('original', encoding='utf-8')
        class BrokenCloseReader(Reader):
            def __exit__(self, *args):
                raise OSError('reader close failed')
        with self.assertRaises(self.backend.ExecutionError):
            self.execute(Client(Instance(BrokenCloseReader())), save=output)
        self.assertEqual(output.read_text(encoding='utf-8'), 'original')
        self.assertEqual(sorted(p.name for p in self.root.iterdir()), ['existing.csv', 'runs'])

    def test_workbook_close_failure_preserves_output_and_cleans_staged_file(self):
        output = self.root / 'existing.xlsx'
        output.write_bytes(b'original')
        with patch('openpyxl.workbook.workbook.Workbook.close', side_effect=OSError('workbook close failed')):
            with self.assertRaises(self.backend.ExecutionError):
                self.execute(save=output)
        self.assertEqual(output.read_bytes(), b'original')
        self.assertEqual(sorted(p.name for p in self.root.iterdir()), ['existing.xlsx', 'runs'])

    def test_final_record_failure_reports_committed_output_as_success_with_warning(self):
        output = self.root / 'existing.csv'
        output.write_text('original', encoding='utf-8')
        real_atomic = self.backend._atomic_text
        def atomic(path, text):
            if path.name == 'state.json' and output.read_text(encoding='utf-8') != 'original':
                raise OSError('record storage unavailable')
            return real_atomic(path, text)
        with patch.object(self.backend, '_atomic_text', side_effect=atomic):
            result, meta = self.execute(save=output)
        self.assertIsNone(result)
        self.assertEqual(meta['state'], 'succeeded')
        self.assertIn('record storage unavailable', meta['record_error'])
        self.assertIn('record storage unavailable', self.stderr.getvalue())
        self.assertEqual(output.read_text(encoding='utf-8-sig').splitlines(), ['value', '0', '1', '2'])

    def test_excel_limits_before_publish_and_actual_constants(self):
        self.assertEqual(self.backend.XLSX_MAX_DATA_ROWS, 1048575)
        self.assertEqual(self.backend.XLSX_MAX_COLUMNS, 16384)
        output = self.root / 'existing.xlsx'
        output.write_bytes(b'original')
        for reader, rows, cols in [(Reader(count=3), 2, 3), (Reader(count=1, columns=('a', 'b')), 3, 1)]:
            with patch.object(self.backend, 'XLSX_MAX_DATA_ROWS', rows), patch.object(self.backend, 'XLSX_MAX_COLUMNS', cols):
                with self.assertRaises(self.backend.ExecutionError) as ctx:
                    self.execute(Client(Instance(reader)), save=output)
            self.assertIn('CSV', str(ctx.exception))
            self.assertEqual(reader.batch_calls, [])
            self.assertEqual(output.read_bytes(), b'original')

    def test_excel_exact_limits_are_allowed(self):
        from openpyxl import load_workbook
        output = self.root / 'boundary.xlsx'
        reader = Reader(count=3, columns=('a', 'b'))
        with patch.object(self.backend, 'XLSX_MAX_DATA_ROWS', 3), patch.object(self.backend, 'XLSX_MAX_COLUMNS', 2):
            _, meta = self.execute(Client(Instance(reader)), save=output, batch_size=2)
        workbook = load_workbook(output, read_only=True)
        try:
            self.assertEqual(list(workbook.active.values), [('a', 'b'), (0, 0), (1, 1), (2, 2)])
        finally:
            workbook.close()
        self.assertEqual(meta['downloaded_rows'], 3)

    def test_xlsx_long_cell_fails_instead_of_silent_truncation(self):
        output = self.root / 'data.xlsx'
        reader = Reader(count=1, batches=[pd.DataFrame({'value': ['x' * 32768]})])
        with self.assertRaises(self.backend.ExecutionError) as ctx:
            self.execute(Client(Instance(reader)), save=output)
        self.assertIn('32767', str(ctx.exception))
        self.assertFalse(output.exists())

    def test_raw_results_use_task_text_with_unknown_row_counts(self):
        for sql in ('desc a', 'describe a', 'show tables', 'explain select 1'):
            client = Client(Instance(query=sql))
            text, meta = self.execute(client, sql=sql)
            self.assertEqual(text, 'raw description')
            self.assertIsNone(meta['total_rows'])
            self.assertEqual(meta['downloaded_rows'], 0)
            self.assertIsNone(meta['truncated'])
            self.assertIsNone(meta['restricted'])
            self.assertEqual(client.instance.reader.preview_calls, [])

    def test_recovered_raw_results_reject_export_without_resubmit(self):
        client = Client(Instance(query='show tables'))
        with self.assertRaises(self.backend.ExecutionError) as ctx:
            self.execute(client, sql=None, instance_id=client.instance.id, save=self.root / 'data.csv')
        self.assertEqual(ctx.exception.exit_code, 6)
        self.assertEqual(client.submissions, [])
        self.assertFalse((self.root / 'data.csv').exists())


if __name__ == '__main__':
    unittest.main()
